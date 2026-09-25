import os
import re
from datetime import datetime
from typing import Dict, Any
import docx
from app.core.config import settings

COMPANY_INFO_REPLACEMENTS = {
    "Địa chỉ: [để trống]": "Địa chỉ: P.910, Tầng 9, Tòa nhà Mapletree Business Centre, 1060 Nguyễn Văn Linh, Phường Tân Hưng, TP Hồ Chí Minh",
    "Mã số thuế: [để trống]": "Mã số thuế: 0301719729",
    "Hotline: [để trống]": "Hotline: 0909 787 040 - 0934 635 766",
    "CÔNG TY TNHH THƯƠNG MẠI DỊCH VỤ PHÚC THANH": "CÔNG TY TNHH THƯƠNG MẠI - DỊCH VỤ PHÚC THÀNH AN",
    "{{SELLER_COMPANY_NAME}}": "CÔNG TY TNHH THƯƠNG MẠI - DỊCH VỤ PHÚC THÀNH AN",
    "{{SELLER_ADDRESS}}": "P.910, Tầng 9, Tòa nhà Mapletree Business Centre, 1060 Nguyễn Văn Linh, Phường Tân Hưng, TP Hồ Chí Minh",
    "{{SELLER_MST}}": "0301719729",
    "{{SELLER_HOTLINE}}": "0909 787 040 - 0934 635 766",
    "{{SELLER_EMAIL}}": "phucthanhaudio@gmail.com",
    "{{SELLER_WEBSITE}}": "phucthanhaudio.vn",
}

def _replace_in_paragraph(paragraph, replacements: Dict[str, str]):
    full_text = paragraph.text
    if not full_text:
        return
    has_match = False
    for k, v in replacements.items():
        if k in full_text:
            full_text = full_text.replace(k, str(v) if v is not None else "")
            has_match = True
    if has_match:
        # Nếu chỉ có 1 run hoặc toàn bộ paragraph có placeholder
        if len(paragraph.runs) > 0:
            paragraph.runs[0].text = full_text
            for r in paragraph.runs[1:]:
                r.text = ""
        else:
            paragraph.text = full_text

def replace_placeholders_in_doc(doc, replacements: Dict[str, str]):
    combined = {**COMPANY_INFO_REPLACEMENTS, **replacements}
    for p in doc.paragraphs:
        _replace_in_paragraph(p, combined)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_in_paragraph(p, combined)

def generate_contract_document(data: Dict[str, Any]) -> str:
    """
    Điền dữ liệu vào template Hợp đồng Word (.docx)
    """
    template_path = os.path.join(settings.INPUT_TEMPLATE_DIR, "HopDong_Mau_PhucThanhAudio_v2.docx")
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Không tìm thấy template Hợp đồng tại: {template_path}")
        
    doc = docx.Document(template_path)
    replace_placeholders_in_doc(doc, data)
    
    contract_id = data.get("{{CONTRACT_ID}}", f"HD_{datetime.now().strftime('%Y%m%d%H%M%S')}")
    filename = f"{contract_id}.docx"
    out_dir = os.path.join(settings.OUTPUT_DIR, "contracts")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, filename)
    doc.save(out_path)
    return out_path

def generate_quote_document(data: Dict[str, Any]) -> str:
    """
    Điền dữ liệu vào template Báo giá ISO Word (.docx)
    """
    template_path = os.path.join(settings.INPUT_TEMPLATE_DIR, "BaoGia_Mau_PhucThanhAudio.docx")
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Không tìm thấy template Báo giá tại: {template_path}")
        
    doc = docx.Document(template_path)
    replace_placeholders_in_doc(doc, data)
    
    quote_id = data.get("{{QUOTE_ID}}", f"BG_{datetime.now().strftime('%Y%m%d%H%M%S')}")
    filename = f"{quote_id}.docx"
    out_dir = os.path.join(settings.OUTPUT_DIR, "quotes")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, filename)
    doc.save(out_path)
    return out_path
